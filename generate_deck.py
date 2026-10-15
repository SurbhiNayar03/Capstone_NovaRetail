from pathlib import Path
import pandas as pd
 
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Pt
 
 
# ---------------------------------------------------------
# File paths
# ---------------------------------------------------------
 
BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "assets"
ICONS_DIR = ASSETS_DIR / "icons"
OUTPUT_DIR = BASE_DIR / "output"
 
DATASET_PATH = ASSETS_DIR / "novaretail_dataset.xlsx"
TEMPLATE_PATH = ASSETS_DIR / "novaretail_template.pptx"
OUTPUT_PATH = OUTPUT_DIR / "NovaRetail_QBR_Q2_2026.pptx"
 
OUTPUT_DIR.mkdir(exist_ok=True)
 
 
# ---------------------------------------------------------
# Load dataset
# ---------------------------------------------------------
 
regional_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Regional_Performance",
    engine="openpyxl",
)
 
customer_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Customer_Experience",
    engine="openpyxl",
)
 
risk_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Operational_Risk",
    engine="openpyxl",
)
 
 
# Clean column names
for dataframe in [regional_df, customer_df, risk_df]:
    dataframe.columns = dataframe.columns.str.strip()
 
 
print("Dataset loaded successfully.")
print(f"Regional performance rows: {len(regional_df)}")
print(f"Customer experience rows: {len(customer_df)}")
print(f"Operational risk rows: {len(risk_df)}")

# ---------------------------------------------------------
# Calculate Q2 2026 KPIs
# ---------------------------------------------------------
 
CURRENT_QUARTER = "2026-Q2"
PREVIOUS_QUARTER = "2026-Q1"
 
 
current_performance = regional_df[
    regional_df["Quarter"] == CURRENT_QUARTER
].copy()
 
previous_performance = regional_df[
    regional_df["Quarter"] == PREVIOUS_QUARTER
].copy()
 
current_customer = customer_df[
    customer_df["Quarter"] == CURRENT_QUARTER
].copy()
 
current_risk = risk_df[
    risk_df["Quarter"] == CURRENT_QUARTER
].copy()
 
 
total_revenue = current_performance["Revenue_USD_M"].sum()
previous_revenue = previous_performance["Revenue_USD_M"].sum()
 
qoq_revenue_growth = (
    (total_revenue - previous_revenue) / previous_revenue
) * 100
 
average_nps = current_customer["NPS"].mean()
 
 
print("\nExecutive KPIs")
print("-" * 40)
print(f"Total Q2 revenue: ${total_revenue:,.0f}M")
print(f"QoQ revenue growth: {qoq_revenue_growth:.1f}%")
print(f"Average regional NPS: {average_nps:.1f}")

# ---------------------------------------------------------
# Regional comparisons
# ---------------------------------------------------------
 
revenue_comparison = current_performance[
    ["Region", "Revenue_USD_M"]
].merge(
    previous_performance[["Region", "Revenue_USD_M"]],
    on="Region",
    suffixes=("_current", "_previous"),
)
 
revenue_comparison["QoQ_Growth_pct"] = (
    (
        revenue_comparison["Revenue_USD_M_current"]
        - revenue_comparison["Revenue_USD_M_previous"]
    )
    / revenue_comparison["Revenue_USD_M_previous"]
) * 100
 
 
strongest_growth = revenue_comparison.loc[
    revenue_comparison["QoQ_Growth_pct"].idxmax()
]
 
weakest_growth = revenue_comparison.loc[
    revenue_comparison["QoQ_Growth_pct"].idxmin()
]
 
best_nps = current_customer.loc[
    current_customer["NPS"].idxmax()
]
 
worst_nps = current_customer.loc[
    current_customer["NPS"].idxmin()
]
 
highest_stockout = current_risk.loc[
    current_risk["Stockout_Rate_pct"].idxmax()
]
 
highest_capacity = current_risk.loc[
    current_risk["Warehouse_Capacity_Util_pct"].idxmax()
]
 
 
print("\nRegional comparisons")
print("-" * 40)
 
print(
    f"Strongest growth: {strongest_growth['Region']} "
    f"({strongest_growth['QoQ_Growth_pct']:.1f}%)"
)
 
print(
    f"Weakest growth: {weakest_growth['Region']} "
    f"({weakest_growth['QoQ_Growth_pct']:.1f}%)"
)
 
print(
    f"Highest NPS: {best_nps['Region']} "
    f"({best_nps['NPS']:.0f})"
)
 
print(
    f"Lowest NPS: {worst_nps['Region']} "
    f"({worst_nps['NPS']:.0f})"
)
 
print(
    f"Highest stockout rate: {highest_stockout['Region']} "
    f"({highest_stockout['Stockout_Rate_pct']:.1f}%)"
)
 
print(
    f"Highest capacity utilization: {highest_capacity['Region']} "
    f"({highest_capacity['Warehouse_Capacity_Util_pct']:.0f}%)"
)

# ---------------------------------------------------------
# Inspect template shapes
# ---------------------------------------------------------
 
def print_template_shapes(presentation):
    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        print(f"\nSLIDE {slide_number}")
        print("-" * 80)
 
        for shape in slide.shapes:
            shape_text = ""
 
            if shape.has_text_frame:
                shape_text = shape.text.replace("\n", " | ")
 
            print(
                f"Name: {shape.name!r} | "
                f"Type: {shape.shape_type} | "
                f"Left: {shape.left} | "
                f"Top: {shape.top} | "
                f"Width: {shape.width} | "
                f"Height: {shape.height} | "
                f"Text: {shape_text[:100]!r}"
            )
 
 
prs = Presentation(TEMPLATE_PATH)
print_template_shapes(prs)

# ---------------------------------------------------------
# Shape helpers
# ---------------------------------------------------------
 
def get_shape_map(slide):
    """
    Return shapes from one slide as:
    {
        "shape_name": shape_object
    }
    """
    return {
        shape.name: shape
        for shape in slide.shapes
    }

def set_shape_text(
    shape,
    text,
    font_size=18,
    bold=False,
    color="1F2937",
    alignment=PP_ALIGN.LEFT,
):
    """
    Replace all existing text in a shape.
    """
    if not shape.has_text_frame:
        raise ValueError(
            f"Shape '{shape.name}' does not contain a text frame."
        )
 
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
 
    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = alignment
 
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
 
 
def set_bullet_text(
    shape,
    bullets,
    font_size=16,
    color="374151",
):
    """
    Replace shape text with formatted bullet points.
    """
    if not shape.has_text_frame:
        raise ValueError(
            f"Shape '{shape.name}' does not contain a text frame."
        )
 
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
 
    for index, bullet in enumerate(bullets):
        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()
 
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(8)

        # ---------------------------------------------------------
# Presentation content
# ---------------------------------------------------------
 
deck_content = {
    "slide_1": {
        "period": "Q2 2026 | Quarterly Business Review",
        "kpi_1_value": f"${total_revenue / 1000:.2f}B",
        "kpi_1_label": "Q2 Revenue",
        "kpi_2_value": f"+{qoq_revenue_growth:.1f}%",
        "kpi_2_label": "QoQ Revenue Growth",
        "kpi_3_value": f"{average_nps:.1f}",
        "kpi_3_label": "Average Regional NPS",
        "summary": (
            "NovaRetail closed Q2 2026 with revenue of "
            f"${total_revenue / 1000:.2f}B, up "
            f"{qoq_revenue_growth:.1f}% from Q1. "
            "APAC delivered the strongest quarterly growth and "
            "the highest customer advocacy score. Latin America "
            "continued to contract while its customer experience "
            "and operational indicators deteriorated. APAC's "
            "capacity constraints also require attention to protect "
            "future growth."
        ),
        "icon": ICONS_DIR / "globe_region.png",
    },
 
    "slide_2": {
        "headline": (
            "APAC Drives Growth as Latin America Declines Further"
        ),
        "bullets": [
            "APAC revenue reached $330M, increasing 8.2% QoQ.",
            "North America remained the largest market at $465M.",
            "Europe improved 4.5% QoQ, but remained below its Q4 peak.",
            "Latin America declined 6.7% QoQ while returns reached 8.2%.",
        ],
        "icon": ICONS_DIR / "growth_chart.png",
    },
 
    "slide_3": {
        "headline": (
            "APAC Sets the CX Benchmark While Latin America Deteriorates"
        ),
        "bullets": [
            "APAC led customer experience with NPS 62 and CSAT of 92%.",
            "North America and Europe improved to NPS 55 and 53.",
            "Latin America NPS fell from 38 to 29 across four quarters.",
            "Delivery reached 7.5 days and support volume climbed to 31K.",
        ],
        "icon": ICONS_DIR / "customer_star.png",
    },
 
    "slide_4": {
        "headline": (
            "Capacity Pressure and Escalating Regional Risk Demand Action"
        ),
        "bullets": [
            "Latin America stockouts reached 12.5%, the highest level.",
            "Its return processing increased to 10 days.",
            "Fraud incidents rose from 18 to 30 across four quarters.",
            "APAC warehouse utilization reached 92%, constraining growth.",
        ],
        "icon": ICONS_DIR / "risk_alert.png",
    },
 
    "slide_5": {
        "bullets": [
            "EXPAND APAC CAPACITY: Improve warehouse capacity and inventory allocation.",
            "STABILIZE LATIN AMERICA: Address fulfillment, returns and support workload.",
            "STRENGTHEN FRAUD CONTROLS: Prioritize preventive transaction monitoring.",
            "REPLICATE CX PRACTICES: Apply APAC delivery practices in weaker markets.",
        ],
        "closing": (
            "Protect APAC's growth while making Latin America's "
            "recovery the immediate leadership priority."
        ),
        "icon": ICONS_DIR / "lightbulb_action.png",
    },
}
 
 # IMPORTANT:
# Replace these example names with the exact names printed
# from your own template.
 

 
from pathlib import Path
 
import pandas as pd
 
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Pt
 
 
# ---------------------------------------------------------
# File paths
# ---------------------------------------------------------
 
BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "assets"
ICONS_DIR = ASSETS_DIR / "icons"
OUTPUT_DIR = BASE_DIR / "output"
 
DATASET_PATH = ASSETS_DIR / "novaretail_dataset.xlsx"
TEMPLATE_PATH = ASSETS_DIR / "novaretail_template.pptx"
OUTPUT_PATH = OUTPUT_DIR / "NovaRetail_QBR_Q2_2026.pptx"
 
OUTPUT_DIR.mkdir(exist_ok=True)
 
 
# ---------------------------------------------------------
# Load dataset
# ---------------------------------------------------------
 
regional_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Regional_Performance",
    engine="openpyxl",
)
 
customer_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Customer_Experience",
    engine="openpyxl",
)
 
risk_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Operational_Risk",
    engine="openpyxl",
)
 
 
# Clean column names
for dataframe in [regional_df, customer_df, risk_df]:
    dataframe.columns = dataframe.columns.str.strip()
 
 
print("Dataset loaded successfully.")
print(f"Regional performance rows: {len(regional_df)}")
print(f"Customer experience rows: {len(customer_df)}")
print(f"Operational risk rows: {len(risk_df)}")

# ---------------------------------------------------------
# Calculate Q2 2026 KPIs
# ---------------------------------------------------------
 
CURRENT_QUARTER = "2026-Q2"
PREVIOUS_QUARTER = "2026-Q1"
 
 
current_performance = regional_df[
    regional_df["Quarter"] == CURRENT_QUARTER
].copy()
 
previous_performance = regional_df[
    regional_df["Quarter"] == PREVIOUS_QUARTER
].copy()
 
current_customer = customer_df[
    customer_df["Quarter"] == CURRENT_QUARTER
].copy()
 
current_risk = risk_df[
    risk_df["Quarter"] == CURRENT_QUARTER
].copy()
 
 
total_revenue = current_performance["Revenue_USD_M"].sum()
previous_revenue = previous_performance["Revenue_USD_M"].sum()
 
qoq_revenue_growth = (
    (total_revenue - previous_revenue) / previous_revenue
) * 100
 
average_nps = current_customer["NPS"].mean()
 
 
print("\nExecutive KPIs")
print("-" * 40)
print(f"Total Q2 revenue: ${total_revenue:,.0f}M")
print(f"QoQ revenue growth: {qoq_revenue_growth:.1f}%")
print(f"Average regional NPS: {average_nps:.1f}")

# ---------------------------------------------------------
# Regional comparisons
# ---------------------------------------------------------
 
revenue_comparison = current_performance[
    ["Region", "Revenue_USD_M"]
].merge(
    previous_performance[["Region", "Revenue_USD_M"]],
    on="Region",
    suffixes=("_current", "_previous"),
)
 
revenue_comparison["QoQ_Growth_pct"] = (
    (
        revenue_comparison["Revenue_USD_M_current"]
        - revenue_comparison["Revenue_USD_M_previous"]
    )
    / revenue_comparison["Revenue_USD_M_previous"]
) * 100
 
 
strongest_growth = revenue_comparison.loc[
    revenue_comparison["QoQ_Growth_pct"].idxmax()
]
 
weakest_growth = revenue_comparison.loc[
    revenue_comparison["QoQ_Growth_pct"].idxmin()
]
 
best_nps = current_customer.loc[
    current_customer["NPS"].idxmax()
]
 
worst_nps = current_customer.loc[
    current_customer["NPS"].idxmin()
]
 
highest_stockout = current_risk.loc[
    current_risk["Stockout_Rate_pct"].idxmax()
]
 
highest_capacity = current_risk.loc[
    current_risk["Warehouse_Capacity_Util_pct"].idxmax()
]
 
 
print("\nRegional comparisons")
print("-" * 40)
 
print(
    f"Strongest growth: {strongest_growth['Region']} "
    f"({strongest_growth['QoQ_Growth_pct']:.1f}%)"
)
 
print(
    f"Weakest growth: {weakest_growth['Region']} "
    f"({weakest_growth['QoQ_Growth_pct']:.1f}%)"
)
 
print(
    f"Highest NPS: {best_nps['Region']} "
    f"({best_nps['NPS']:.0f})"
)
 
print(
    f"Lowest NPS: {worst_nps['Region']} "
    f"({worst_nps['NPS']:.0f})"
)
 
print(
    f"Highest stockout rate: {highest_stockout['Region']} "
    f"({highest_stockout['Stockout_Rate_pct']:.1f}%)"
)
 
print(
    f"Highest capacity utilization: {highest_capacity['Region']} "
    f"({highest_capacity['Warehouse_Capacity_Util_pct']:.0f}%)"
)

# ---------------------------------------------------------
# Inspect template shapes
# ---------------------------------------------------------
 
def print_template_shapes(presentation):
    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        print(f"\nSLIDE {slide_number}")
        print("-" * 80)
 
        for shape in slide.shapes:
            shape_text = ""
 
            if shape.has_text_frame:
                shape_text = shape.text.replace("\n", " | ")
 
            print(
                f"Name: {shape.name!r} | "
                f"Type: {shape.shape_type} | "
                f"Left: {shape.left} | "
                f"Top: {shape.top} | "
                f"Width: {shape.width} | "
                f"Height: {shape.height} | "
                f"Text: {shape_text[:100]!r}"
            )
 
 
prs = Presentation(TEMPLATE_PATH)
print_template_shapes(prs)

# ---------------------------------------------------------
# Shape helpers
# ---------------------------------------------------------
 
def get_shape_map(slide):
    """
    Return shapes from one slide as:
    {
        "shape_name": shape_object
    }
    """
    return {
        shape.name: shape
        for shape in slide.shapes
    }

def set_shape_text(
    shape,
    text,
    font_size=18,
    bold=False,
    color="1F2937",
    alignment=PP_ALIGN.LEFT,
):
    """
    Replace all existing text in a shape.
    """
    if not shape.has_text_frame:
        raise ValueError(
            f"Shape '{shape.name}' does not contain a text frame."
        )
 
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
 
    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = alignment
 
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
 
 
def set_bullet_text(
    shape,
    bullets,
    font_size=16,
    color="374151",
):
    """
    Replace shape text with formatted bullet points.
    """
    if not shape.has_text_frame:
        raise ValueError(
            f"Shape '{shape.name}' does not contain a text frame."
        )
 
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
 
    for index, bullet in enumerate(bullets):
        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()
 
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(8)

        # ---------------------------------------------------------
# Presentation content
# ---------------------------------------------------------
 
deck_content = {
    "slide_1": {
        "period": "Q2 2026 | Quarterly Business Review",
        "kpi_1_value": f"${total_revenue / 1000:.2f}B",
        "kpi_1_label": "Q2 Revenue",
        "kpi_2_value": f"+{qoq_revenue_growth:.1f}%",
        "kpi_2_label": "QoQ Revenue Growth",
        "kpi_3_value": f"{average_nps:.1f}",
        "kpi_3_label": "Average Regional NPS",
        "summary": (
            "NovaRetail closed Q2 2026 with revenue of "
            f"${total_revenue / 1000:.2f}B, up "
            f"{qoq_revenue_growth:.1f}% from Q1. "
            "APAC delivered the strongest quarterly growth and "
            "the highest customer advocacy score. Latin America "
            "continued to contract while its customer experience "
            "and operational indicators deteriorated. APAC's "
            "capacity constraints also require attention to protect "
            "future growth."
        ),
        "icon": ICONS_DIR / "globe_region.png",
    },
 
    "slide_2": {
        "headline": (
            "APAC Drives Growth as Latin America Declines Further"
        ),
        "bullets": [
            "APAC revenue reached $330M, increasing 8.2% QoQ.",
            "North America remained the largest market at $465M.",
            "Europe improved 4.5% QoQ, but remained below its Q4 peak.",
            "Latin America declined 6.7% QoQ while returns reached 8.2%.",
        ],
        "icon": ICONS_DIR / "growth_chart.png",
    },
 
    "slide_3": {
        "headline": (
            "APAC Sets the CX Benchmark While Latin America Deteriorates"
        ),
        "bullets": [
            "APAC led customer experience with NPS 62 and CSAT of 92%.",
            "North America and Europe improved to NPS 55 and 53.",
            "Latin America NPS fell from 38 to 29 across four quarters.",
            "Delivery reached 7.5 days and support volume climbed to 31K.",
        ],
        "icon": ICONS_DIR / "customer_star.png",
    },
 
    "slide_4": {
        "headline": (
            "Capacity Pressure and Escalating Regional Risk Demand Action"
        ),
        "bullets": [
            "Latin America stockouts reached 12.5%, the highest level.",
            "Its return processing increased to 10 days.",
            "Fraud incidents rose from 18 to 30 across four quarters.",
            "APAC warehouse utilization reached 92%, constraining growth.",
        ],
        "icon": ICONS_DIR / "risk_alert.png",
    },
 
    "slide_5": {
        "bullets": [
            "EXPAND APAC CAPACITY: Improve warehouse capacity and inventory allocation.",
            "STABILIZE LATIN AMERICA: Address fulfillment, returns and support workload.",
            "STRENGTHEN FRAUD CONTROLS: Prioritize preventive transaction monitoring.",
            "REPLICATE CX PRACTICES: Apply APAC delivery practices in weaker markets.",
        ],
        "closing": (
            "Protect APAC's growth while making Latin America's "
            "recovery the immediate leadership priority."
        ),
        "icon": ICONS_DIR / "lightbulb_action.png",
    },
}
 
 # IMPORTANT:
# Replace these example names with the exact names printed
# from your own template.
 
SHAPE_NAMES = {
    1: {
        "period": "subtitle_1",
        "kpi_1_value": "kpi1_value",
        "kpi_1_label": "kpi1_label",
        "kpi_2_value": "kpi2_value",
        "kpi_2_label": "kpi2_label",
        "kpi_3_value": "kpi3_value",
        "kpi_3_label": "kpi3_label",
        "summary": "summary_text",
        "icon": "icon_1",
    },
 
    2: {
        "headline": "title_2",
        "chart": "chart_performance",
        "insights": "insights_performance",
        "icon": "icon_2",
    },
 
    3: {
        "headline": "title_3",
        "chart": "chart_customer",
        "insights": "insights_customer",
        "icon": "icon_3",
    },
 
    4: {
        "headline": "title_4",
        "chart": "chart_risk",
        "insights": "insights_risk",
        "icon": "icon_4",
    },
 
    5: {
        "actions": "recommendations_text",
        "closing": "closing_statement",
        "icon": "icon_5",
    },
}
 

def get_named_shape(slide, shape_name):
    shapes = get_shape_map(slide)
 
    if shape_name not in shapes:
        available_shapes = "\n".join(
            f" - {name}"
            for name in shapes
        )
 
        raise KeyError(
            f"\nShape '{shape_name}' was not found.\n"
            f"Available shapes are:\n{available_shapes}"
        )
 
    return shapes[shape_name]

def capture_placeholder_position(shape):
    """
    Save the placeholder position before deleting it.
    """
    return {
        "left": shape.left,
        "top": shape.top,
        "width": shape.width,
        "height": shape.height,
    }
 
 
def remove_shape(shape):
    """
    Remove a shape from the slide's XML tree.
    """
    shape_element = shape._element
    shape_element.getparent().remove(shape_element)

def add_performance_chart(slide, placeholder_shape, dataframe):
    position = capture_placeholder_position(
        placeholder_shape
    )
 
    remove_shape(placeholder_shape)
 
    quarters = dataframe["Quarter"].drop_duplicates().tolist()
    regions = dataframe["Region"].drop_duplicates().tolist()
 
    chart_data = CategoryChartData()
    chart_data.categories = regions
 
    for quarter in quarters:
        quarter_data = (
            dataframe[dataframe["Quarter"] == quarter]
            .set_index("Region")
            .reindex(regions)
        )
 
        chart_data.add_series(
            quarter,
            tuple(quarter_data["Revenue_USD_M"]),
        )
 
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        position["left"],
        position["top"],
        position["width"],
        position["height"],
        chart_data,
    )
 
    chart = graphic_frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Revenue by Region ($M)"
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
 
    return graphic_frame

def add_customer_chart(slide, placeholder_shape, dataframe):
    position = capture_placeholder_position(
        placeholder_shape
    )
 
    remove_shape(placeholder_shape)
 
    quarters = dataframe["Quarter"].drop_duplicates().tolist()
    regions = dataframe["Region"].drop_duplicates().tolist()
 
    chart_data = CategoryChartData()
    chart_data.categories = quarters
 
    for region in regions:
        region_data = (
            dataframe[dataframe["Region"] == region]
            .set_index("Quarter")
            .reindex(quarters)
        )
 
        chart_data.add_series(
            region,
            tuple(region_data["NPS"]),
        )
 
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        position["left"],
        position["top"],
        position["width"],
        position["height"],
        chart_data,
    )
 
    chart = graphic_frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "NPS Trend by Region"
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 70
 
    return graphic_frame

def add_risk_chart(slide, placeholder_shape, dataframe):
    position = capture_placeholder_position(
        placeholder_shape
    )
 
    remove_shape(placeholder_shape)
 
    q2_risk = dataframe[
        dataframe["Quarter"] == CURRENT_QUARTER
    ].copy()
 
    regions = q2_risk["Region"].tolist()
 
    chart_data = CategoryChartData()
    chart_data.categories = regions
 
    chart_data.add_series(
        "Stockout Rate %",
        tuple(q2_risk["Stockout_Rate_pct"]),
    )
 
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        position["left"],
        position["top"],
        position["width"],
        position["height"],
        chart_data,
    )
 
    chart = graphic_frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Q2 2026 Stockout Rate"
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 15
 
    return graphic_frame

def replace_placeholder_with_picture(
    slide,
    placeholder_shape,
    image_path,
):
    if not image_path.exists():
        raise FileNotFoundError(
            f"Icon was not found: {image_path}"
        )
 
    position = capture_placeholder_position(
        placeholder_shape
    )
 
    remove_shape(placeholder_shape)
 
    picture = slide.shapes.add_picture(
        str(image_path),
        position["left"],
        position["top"],
        position["width"],
        position["height"],
    )
 
    return picture

def populate_slide_1(prs):
    slide = prs.slides[0]
    names = SHAPE_NAMES[1]
    content = deck_content["slide_1"]
 
    set_shape_text(
        get_named_shape(slide, names["period"]),
        content["period"],
        font_size=16,
        bold=True,
        color="4B5563",
    )
 
    set_shape_text(
        get_named_shape(slide, names["kpi_1_value"]),
        content["kpi_1_value"],
        font_size=26,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )
 
    set_shape_text(
        get_named_shape(slide, names["kpi_1_label"]),
        content["kpi_1_label"],
        font_size=11,
        color="4B5563",
        alignment=PP_ALIGN.CENTER,
    )
 
    set_shape_text(
        get_named_shape(slide, names["kpi_2_value"]),
        content["kpi_2_value"],
        font_size=26,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )
 
    set_shape_text(
        get_named_shape(slide, names["kpi_2_label"]),
        content["kpi_2_label"],
        font_size=11,
        color="4B5563",
        alignment=PP_ALIGN.CENTER,
    )
 
    set_shape_text(
        get_named_shape(slide, names["kpi_3_value"]),
        content["kpi_3_value"],
        font_size=26,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )
 
    set_shape_text(
        get_named_shape(slide, names["kpi_3_label"]),
        content["kpi_3_label"],
        font_size=11,
        color="4B5563",
        alignment=PP_ALIGN.CENTER,
    )
 
    set_shape_text(
        get_named_shape(slide, names["summary"]),
        content["summary"],
        font_size=15,
        color="374151",
    )
 
    replace_placeholder_with_picture(
        slide,
        get_named_shape(slide, names["icon"]),
        content["icon"],
    )
 
def populate_slide_2(prs):
    slide = prs.slides[1]
    names = SHAPE_NAMES[2]
    content = deck_content["slide_2"]
 
    set_shape_text(
        get_named_shape(slide, names["headline"]),
        content["headline"],
        font_size=24,
        bold=True,
        color="111827",
    )
 
    set_bullet_text(
        get_named_shape(slide, names["insights"]),
        content["bullets"],
        font_size=14,
    )
 
    chart_placeholder = get_named_shape(
        slide,
        names["chart"],
    )
 
    add_performance_chart(
        slide,
        chart_placeholder,
        regional_df,
    )
 
    replace_placeholder_with_picture(
        slide,
        get_named_shape(slide, names["icon"]),
        content["icon"],
    )
 
 
def populate_slide_3(prs):
    slide = prs.slides[2]
    names = SHAPE_NAMES[3]
    content = deck_content["slide_3"]
 
    set_shape_text(
        get_named_shape(slide, names["headline"]),
        content["headline"],
        font_size=24,
        bold=True,
        color="111827",
    )
 
    set_bullet_text(
        get_named_shape(slide, names["insights"]),
        content["bullets"],
        font_size=14,
    )
 
    chart_placeholder = get_named_shape(
        slide,
        names["chart"],
    )
 
    add_customer_chart(
        slide,
        chart_placeholder,
        customer_df,
    )
 
    replace_placeholder_with_picture(
        slide,
        get_named_shape(slide, names["icon"]),
        content["icon"],
    )
 
 
def populate_slide_4(prs):
    slide = prs.slides[3]
    names = SHAPE_NAMES[4]
    content = deck_content["slide_4"]
 
    set_shape_text(
        get_named_shape(slide, names["headline"]),
        content["headline"],
        font_size=24,
        bold=True,
        color="111827",
    )
 
    set_bullet_text(
        get_named_shape(slide, names["insights"]),
        content["bullets"],
        font_size=14,
    )
 
    chart_placeholder = get_named_shape(
        slide,
        names["chart"],
    )
 
    add_risk_chart(
        slide,
        chart_placeholder,
        risk_df,
    )
 
    replace_placeholder_with_picture(
        slide,
        get_named_shape(slide, names["icon"]),
        content["icon"],
    )

def set_action_bullets(shape, actions):
    if not shape.has_text_frame:
        raise ValueError(
            f"Shape '{shape.name}' has no text frame."
        )
 
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
 
    for index, action in enumerate(actions):
        heading, description = action.split(":", maxsplit=1)
 
        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()
 
        paragraph.level = 0
        paragraph.space_after = Pt(10)
 
        heading_run = paragraph.add_run()
        heading_run.text = heading + ":"
        heading_run.font.name = "Aptos"
        heading_run.font.size = Pt(15)
        heading_run.font.bold = True
        heading_run.font.color.rgb = RGBColor.from_string(
            "0F766E"
        )
 
        description_run = paragraph.add_run()
        description_run.text = description
        description_run.font.name = "Aptos"
        description_run.font.size = Pt(15)
        description_run.font.color.rgb = RGBColor.from_string(
            "374151"
        )
 
 
def populate_slide_5(prs):
    slide = prs.slides[4]
    names = SHAPE_NAMES[5]
    content = deck_content["slide_5"]
 
    set_action_bullets(
        get_named_shape(slide, names["actions"]),
        content["bullets"],
    )
 
    set_shape_text(
        get_named_shape(slide, names["closing"]),
        content["closing"],
        font_size=18,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )
 
    replace_placeholder_with_picture(
        slide,
        get_named_shape(slide, names["icon"]),
        content["icon"],
    )
 
def remove_remaining_instructions(prs):
    instruction_markers = [
        "[FILL:",
        "[VALUE]",
        "Insert your chart here",
        "Pick the icon",
    ]
 
    for slide in prs.slides:
        for shape in list(slide.shapes):
            if not shape.has_text_frame:
                continue
 
            text = shape.text.strip()
 
            contains_instruction = any(
                marker.lower() in text.lower()
                for marker in instruction_markers
            )
 
            if contains_instruction:
                shape.text_frame.clear()

# ---------------------------------------------------------
# Main workflow
# ---------------------------------------------------------
 
def main():
    if not DATASET_PATH.exists():
        raise FileNotFoundError(
            f"Dataset not found: {DATASET_PATH}"
        )
 
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Template not found: {TEMPLATE_PATH}"
        )
 
    prs = Presentation(TEMPLATE_PATH)
 
    if len(prs.slides) != 5:
        raise ValueError(
            f"Expected 5 slides, found {len(prs.slides)}."
        )
 
    populate_slide_1(prs)
    populate_slide_2(prs)
    populate_slide_3(prs)
    populate_slide_4(prs)
    populate_slide_5(prs)
 
    remove_remaining_instructions(prs)
 
    prs.save(OUTPUT_PATH)
 
    print("\nPresentation generated successfully:")
    print(OUTPUT_PATH)
 
 
if __name__ == "__main__":
    main()





      


 